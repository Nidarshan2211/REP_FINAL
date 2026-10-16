import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # =========================================================================
    # SLIDE 1: Resident Dashboard Alert View (Matches User's Slide 19 Design)
    # =========================================================================
    slide1 = prs.slides.add_slide(prs.slide_layouts[6])

    # Background
    bg1 = slide1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg1.line.fill.background()

    # Title & Subtitle
    title_box = slide1.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(7.5), Inches(0.55))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = "Explainable Alerts: Reason Codes & Severity"
    p.font.name = "Georgia"
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = RGBColor(30, 41, 59)

    sub_box = slide1.shapes.add_textbox(Inches(0.8), Inches(0.98), Inches(7.5), Inches(0.35))
    stf = sub_box.text_frame
    stf.word_wrap = True
    stf.margin_left = stf.margin_right = stf.margin_top = stf.margin_bottom = 0
    sp = stf.paragraphs[0]
    sp.text = "Residents see why an alert fired, not just that it fired"
    sp.font.name = "Segoe UI"
    sp.font.size = Pt(13)
    sp.font.color.rgb = RGBColor(100, 116, 139)

    # Severity Levels Section
    sev_title = slide1.shapes.add_textbox(Inches(8.5), Inches(0.35), Inches(4.0), Inches(0.25))
    stf2 = sev_title.text_frame
    stf2.margin_left = stf2.margin_right = stf2.margin_top = stf2.margin_bottom = 0
    sp2 = stf2.paragraphs[0]
    sp2.alignment = PP_ALIGN.RIGHT
    sp2.text = "SEVERITY LEVELS"
    sp2.font.name = "Segoe UI"
    sp2.font.size = Pt(11)
    sp2.font.bold = True
    sp2.font.color.rgb = RGBColor(217, 119, 6)

    levels = [
        ("LOW", RGBColor(245, 158, 11)),
        ("MEDIUM", RGBColor(234, 88, 12)),
        ("HIGH", RGBColor(220, 38, 38)),
        ("CRITICAL", RGBColor(69, 26, 3))
    ]
    badge_w = Inches(0.95)
    badge_h = Inches(0.55)
    start_bx = Inches(8.55)
    spacing_badge = Inches(0.08)

    for i, (lvl, color) in enumerate(levels):
        bx = start_bx + i * (badge_w + spacing_badge)
        b_shape = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, bx, Inches(0.65), badge_w, badge_h)
        b_shape.fill.solid()
        b_shape.fill.fore_color.rgb = color
        b_shape.line.fill.background()
        btf = b_shape.text_frame
        btf.margin_left = btf.margin_right = btf.margin_top = btf.margin_bottom = 0
        bp = btf.paragraphs[0]
        bp.alignment = PP_ALIGN.CENTER
        bp.text = lvl
        bp.font.name = "Segoe UI"
        bp.font.size = Pt(10)
        bp.font.bold = True
        bp.font.color.rgb = RGBColor(255, 255, 255)

    # Reason Card 1: Abnormal Spike
    card_y = Inches(1.50)
    card_h = Inches(1.15)
    card_w = Inches(5.72)
    card1_x = Inches(0.8)
    card2_x = Inches(6.81)

    rc1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card1_x, card_y, card_w, card_h)
    rc1.fill.solid()
    rc1.fill.fore_color.rgb = RGBColor(255, 247, 237)
    rc1.line.color.rgb = RGBColor(254, 215, 170)
    rc1.line.width = Pt(1)

    ic1 = slide1.shapes.add_shape(MSO_SHAPE.OVAL, card1_x + Inches(0.20), card_y + Inches(0.25), Inches(0.65), Inches(0.65))
    ic1.fill.solid()
    ic1.fill.fore_color.rgb = RGBColor(234, 88, 12)
    ic1.line.fill.background()

    t1_box = slide1.shapes.add_textbox(card1_x + Inches(0.95), card_y + Inches(0.18), card_w - Inches(1.1), Inches(0.80))
    t1_tf = t1_box.text_frame
    t1_tf.word_wrap = True
    t1_tf.margin_left = t1_tf.margin_right = t1_tf.margin_top = t1_tf.margin_bottom = 0
    t1_p1 = t1_tf.paragraphs[0]
    t1_p1.text = "ABNORMAL_SPIKE_FOR_HOUSEHOLD_SIZE"
    t1_p1.font.name = "Segoe UI"
    t1_p1.font.size = Pt(11)
    t1_p1.font.bold = True
    t1_p1.font.color.rgb = RGBColor(194, 65, 12)

    t1_p2 = t1_tf.add_paragraph()
    t1_p2.text = "Usage far above what the declared household size explains (e.g. 48.5 kL vs 16.8 kL upper ceiling threshold for 4 members)."
    t1_p2.font.name = "Segoe UI"
    t1_p2.font.size = Pt(9.5)
    t1_p2.font.color.rgb = RGBColor(71, 85, 105)

    # Reason Card 2: Away Consumption
    rc2 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, card2_x, card_y, card_w, card_h)
    rc2.fill.solid()
    rc2.fill.fore_color.rgb = RGBColor(255, 247, 237)
    rc2.line.color.rgb = RGBColor(254, 215, 170)
    rc2.line.width = Pt(1)

    ic2 = slide1.shapes.add_shape(MSO_SHAPE.OVAL, card2_x + Inches(0.20), card_y + Inches(0.25), Inches(0.65), Inches(0.65))
    ic2.fill.solid()
    ic2.fill.fore_color.rgb = RGBColor(220, 38, 38)
    ic2.line.fill.background()

    t2_box = slide1.shapes.add_textbox(card2_x + Inches(0.95), card_y + Inches(0.18), card_w - Inches(1.1), Inches(0.80))
    t2_tf = t2_box.text_frame
    t2_tf.word_wrap = True
    t2_tf.margin_left = t2_tf.margin_right = t2_tf.margin_top = t2_tf.margin_bottom = 0
    t2_p1 = t2_tf.paragraphs[0]
    t2_p1.text = "UNEXPLAINED_AWAY_CONSUMPTION"
    t2_p1.font.name = "Segoe UI"
    t2_p1.font.size = Pt(11)
    t2_p1.font.bold = True
    t2_p1.font.color.rgb = RGBColor(185, 28, 28)

    t2_p2 = t2_tf.add_paragraph()
    t2_p2.text = "Significant water consumption detected while flat occupancy is marked as vacant or resident logged an active vacation period."
    t2_p2.font.name = "Segoe UI"
    t2_p2.font.size = Pt(9.5)
    t2_p2.font.color.rgb = RGBColor(71, 85, 105)

    # Screenshot Section
    img1_path = "Cropped_Resident_Dashboard_Alert.png"
    if os.path.exists(img1_path):
        img_x = Inches(0.8)
        img_y = Inches(2.80)
        img_w = Inches(11.733)
        img_h = Inches(4.10)

        frame1 = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, img_x, img_y, img_w, img_h)
        frame1.fill.solid()
        frame1.fill.fore_color.rgb = RGBColor(255, 255, 255)
        frame1.line.color.rgb = RGBColor(226, 232, 240)
        frame1.line.width = Pt(1.5)

        slide1.shapes.add_picture(img1_path, img_x + Inches(0.06), img_y + Inches(0.06), width=img_w - Inches(0.12), height=img_h - Inches(0.12))

    # Footer 1
    foot1 = slide1.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(10.0), Inches(0.30))
    ftf1 = foot1.text_frame
    ftf1.margin_left = ftf1.margin_right = ftf1.margin_top = ftf1.margin_bottom = 0
    fp1 = ftf1.paragraphs[0]
    fp1.text = "AQATRAC | Web-Based Water Billing & Leak Management"
    fp1.font.name = "Segoe UI"
    fp1.font.size = Pt(9)
    fp1.font.color.rgb = RGBColor(148, 163, 184)

    page1 = slide1.shapes.add_textbox(Inches(11.5), Inches(7.05), Inches(1.0), Inches(0.30))
    ptf1 = page1.text_frame
    ptf1.margin_left = ptf1.margin_right = ptf1.margin_top = ptf1.margin_bottom = 0
    pp1 = ptf1.paragraphs[0]
    pp1.alignment = PP_ALIGN.RIGHT
    pp1.text = "19"
    pp1.font.name = "Segoe UI"
    pp1.font.size = Pt(9.5)
    pp1.font.color.rgb = RGBColor(148, 163, 184)


    # =========================================================================
    # SLIDE 2: Dedicated Diagnostics & Active Feedback View (Companion Slide)
    # =========================================================================
    slide2 = prs.slides.add_slide(prs.slide_layouts[6])

    bg2 = slide2.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg2.fill.solid()
    bg2.fill.fore_color.rgb = RGBColor(255, 255, 255)
    bg2.line.fill.background()

    # Title & Subtitle 2
    t2_box = slide2.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(8.5), Inches(0.55))
    tf2 = t2_box.text_frame
    tf2.word_wrap = True
    tf2.margin_left = tf2.margin_right = tf2.margin_top = tf2.margin_bottom = 0
    p2 = tf2.paragraphs[0]
    p2.text = "Leak Diagnostics: Risk Assessment & Feedback Actions"
    p2.font.name = "Georgia"
    p2.font.size = Pt(24)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(30, 41, 59)

    sub2_box = slide2.shapes.add_textbox(Inches(0.8), Inches(0.98), Inches(8.5), Inches(0.35))
    stf2 = sub2_box.text_frame
    stf2.word_wrap = True
    stf2.margin_left = stf2.margin_right = stf2.margin_top = stf2.margin_bottom = 0
    sp2 = stf2.paragraphs[0]
    sp2.text = "Residents inspect calibrated risk bands, estimated extra costs, and log feedback for AI model retraining"
    sp2.font.name = "Segoe UI"
    sp2.font.size = Pt(13)
    sp2.font.color.rgb = RGBColor(100, 116, 139)

    # Screenshot Section 2: Diagnostics Hero Card
    img2_path = "Cropped_Resident_Diagnostics_Hero.png"
    if os.path.exists(img2_path):
        img_x = Inches(0.8)
        img_y = Inches(1.50)
        img_w = Inches(11.733)
        img_h = Inches(5.35)

        frame2 = slide2.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, img_x, img_y, img_w, img_h)
        frame2.fill.solid()
        frame2.fill.fore_color.rgb = RGBColor(255, 255, 255)
        frame2.line.color.rgb = RGBColor(226, 232, 240)
        frame2.line.width = Pt(1.5)

        slide2.shapes.add_picture(img2_path, img_x + Inches(0.06), img_y + Inches(0.06), width=img_w - Inches(0.12), height=img_h - Inches(0.12))

    # Footer 2
    foot2 = slide2.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(10.0), Inches(0.30))
    ftf2 = foot2.text_frame
    ftf2.margin_left = ftf2.margin_right = ftf2.margin_top = ftf2.margin_bottom = 0
    fp2 = ftf2.paragraphs[0]
    fp2.text = "AQATRAC | Web-Based Water Billing & Leak Management"
    fp2.font.name = "Segoe UI"
    fp2.font.size = Pt(9)
    fp2.font.color.rgb = RGBColor(148, 163, 184)

    page2 = slide2.shapes.add_textbox(Inches(11.5), Inches(7.05), Inches(1.0), Inches(0.30))
    ptf2 = page2.text_frame
    ptf2.margin_left = ptf2.margin_right = ptf2.margin_top = ptf2.margin_bottom = 0
    pp2 = ptf2.paragraphs[0]
    pp2.alignment = PP_ALIGN.RIGHT
    pp2.text = "20"
    pp2.font.name = "Segoe UI"
    pp2.font.size = Pt(9.5)
    pp2.font.color.rgb = RGBColor(148, 163, 184)

    out_file = "AquaFlow_Explainable_Alerts_Slide.pptx"
    prs.save(out_file)
    print(f"Presentation saved to: {out_file}")

if __name__ == "__main__":
    build_presentation()
