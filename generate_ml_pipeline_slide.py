import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_ml_pipeline_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)

    # Background: Crisp deep navy header banner with clean slate background
    bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = RGBColor(248, 250, 252) # Slate-50
    bg.line.color.rgb = RGBColor(248, 250, 252)

    # Top Brand Header Bar
    top_bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(1.35))
    top_bar.fill.solid()
    top_bar.fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate-900 (Navy)
    top_bar.line.color.rgb = RGBColor(15, 23, 42)

    # Accent Strip
    accent = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, Inches(1.30), Inches(13.333), Inches(0.05))
    accent.fill.solid()
    accent.fill.fore_color.rgb = RGBColor(14, 165, 233) # Sky-500
    accent.line.color.rgb = RGBColor(14, 165, 233)

    # Title & Subtitle in Header
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.18), Inches(11.7), Inches(0.55))
    tf = title_box.text_frame
    tf.word_wrap = True
    tf.margin_top = tf.margin_bottom = tf.margin_left = tf.margin_right = 0
    p = tf.paragraphs[0]
    p.text = "MACHINE LEARNING LEAK DETECTION PIPELINE"
    p.font.name = "Segoe UI"
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)

    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(11.7), Inches(0.40))
    stf = sub_box.text_frame
    stf.word_wrap = True
    stf.margin_top = stf.margin_bottom = stf.margin_left = stf.margin_right = 0
    sp = stf.paragraphs[0]
    sp.text = "End-to-End Architecture: Monthly Telemetry Normalization, LightGBM Quantile Regression & Continuous Risk Attribution"
    sp.font.name = "Segoe UI"
    sp.font.size = Pt(11)
    sp.font.color.rgb = RGBColor(148, 163, 184) # Slate-400

    # 4 Pipeline Stages (Cards)
    stages = [
        {
            "num": "STEP 1",
            "title": "Data Ingestion & Normalization",
            "icon": "📥",
            "color": RGBColor(2, 132, 199), # Sky Blue
            "items": [
                ("Monthly Meter Reading", "Entered in kL per billing cycle (e.g. 48.5 kL)"),
                ("Cycle Day Normalization", "daily_usage = total_kL / days_in_cycle (30/31d)"),
                ("Household Demographics", "Tenant count, BHK unit type, floor level"),
                ("Away Days Deduction", "Active occupancy days adjusted for vacations")
            ]
        },
        {
            "num": "STEP 2",
            "title": "21-D Feature Engineering",
            "icon": "⚙️",
            "color": RGBColor(99, 102, 241), # Indigo
            "items": [
                ("Temporal Lags", "lag_1, lag_2, lag_3 (past 3 cycles daily consumption)"),
                ("Rolling Stats", "6-month rolling mean, median, standard deviation"),
                ("Trend Slope", "Linear regression trajectory across previous cycles"),
                ("Seasonality & Peers", "Cyclical sin/cos(month) + Building BHK peer median")
            ]
        },
        {
            "num": "STEP 3",
            "title": "Quantile Regression (ML)",
            "icon": "🧠",
            "color": RGBColor(16, 185, 129), # Emerald Green
            "items": [
                ("Algorithm", "LightGBM Quantile Regressors (Pinball Loss)"),
                ("Q10 (Lower Bound)", "Conservative minimum expected baseline usage"),
                ("Q50 (Median Baseline)", "Most probable expected normal monthly volume"),
                ("Q95 (Upper Ceiling)", "Statistical threshold; exceedance triggers leak check")
            ]
        },
        {
            "num": "STEP 4",
            "title": "Risk Scoring & Diagnostics",
            "icon": "🎯",
            "color": RGBColor(239, 68, 68), # Coral Red
            "items": [
                ("Continuous Risk Score", "Scaled 0.0 to 1.0 (displayed as 0% to 100% in UI)"),
                ("Status Classification", "NORMAL (<40%), WATCH (40-69%), POTENTIAL_LEAK (>=70%)"),
                ("Pattern Detection", "SUDDEN_SPIKE, GRADUAL_DRIFT, STEP_CHANGE"),
                ("Explainable Diagnosis", "Plain English reason + Extra Cost (₹) + Feedback loop")
            ]
        }
    ]

    card_width = Inches(2.75)
    card_height = Inches(3.90)
    card_y = Inches(1.55)
    spacing = Inches(0.24)
    start_x = Inches(0.80)

    for i, s in enumerate(stages):
        cx = start_x + i * (card_width + spacing)

        # Card Container (Rounded Rectangle)
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx, card_y, card_width, card_height)
        card.fill.solid()
        card.fill.fore_color.rgb = RGBColor(255, 255, 255)
        card.line.color.rgb = RGBColor(226, 232, 240) # Slate-200
        card.line.width = Pt(1.5)

        # Step Header Pill
        pill = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, cx + Inches(0.15), card_y + Inches(0.15), Inches(1.1), Inches(0.32))
        pill.fill.solid()
        pill.fill.fore_color.rgb = s["color"]
        pill.line.color.rgb = s["color"]
        ptf = pill.text_frame
        ptf.margin_top = ptf.margin_bottom = ptf.margin_left = ptf.margin_right = 0
        pp = ptf.paragraphs[0]
        pp.alignment = PP_ALIGN.CENTER
        pp.text = s["num"]
        pp.font.name = "Segoe UI"
        pp.font.size = Pt(9.5)
        pp.font.bold = True
        pp.font.color.rgb = RGBColor(255, 255, 255)

        # Title
        t_box = slide.shapes.add_textbox(cx + Inches(0.15), card_y + Inches(0.55), card_width - Inches(0.30), Inches(0.55))
        ttf = t_box.text_frame
        ttf.word_wrap = True
        ttf.margin_top = ttf.margin_bottom = ttf.margin_left = ttf.margin_right = 0
        tp = ttf.paragraphs[0]
        tp.text = s["title"]
        tp.font.name = "Segoe UI"
        tp.font.size = Pt(12)
        tp.font.bold = True
        tp.font.color.rgb = RGBColor(15, 23, 42)

        # Divider line
        div = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, cx + Inches(0.15), card_y + Inches(1.15), card_width - Inches(0.30), Inches(0.015))
        div.fill.solid()
        div.fill.fore_color.rgb = RGBColor(241, 245, 249)
        div.line.color.rgb = RGBColor(241, 245, 249)

        # Items
        item_y = card_y + Inches(1.25)
        for label, desc in s["items"]:
            ibox = slide.shapes.add_textbox(cx + Inches(0.15), item_y, card_width - Inches(0.30), Inches(0.60))
            itf = ibox.text_frame
            itf.word_wrap = True
            itf.margin_top = itf.margin_bottom = itf.margin_left = itf.margin_right = 0
            
            p1 = itf.paragraphs[0]
            p1.text = "• " + label
            p1.font.name = "Segoe UI"
            p1.font.size = Pt(10)
            p1.font.bold = True
            p1.font.color.rgb = s["color"]

            p2 = itf.add_paragraph()
            p2.text = desc
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(9)
            p2.font.color.rgb = RGBColor(71, 85, 105) # Slate-600

            item_y += Inches(0.62)

    # Bottom Banner: Real Demonstration Walkthrough Case Study
    bot_y = Inches(5.65)
    bot_w = Inches(11.733)
    bot_h = Inches(1.55)

    bot_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.80), bot_y, bot_w, bot_h)
    bot_card.fill.solid()
    bot_card.fill.fore_color.rgb = RGBColor(15, 23, 42) # Slate-900
    bot_card.line.color.rgb = RGBColor(30, 41, 59)
    bot_card.line.width = Pt(1)

    # Bottom Content
    bbox = slide.shapes.add_textbox(Inches(1.0), bot_y + Inches(0.15), bot_w - Inches(0.4), bot_h - Inches(0.3))
    btf = bbox.text_frame
    btf.word_wrap = True
    btf.margin_top = btf.margin_bottom = btf.margin_left = btf.margin_right = 0

    bp1 = btf.paragraphs[0]
    bp1.text = "LIVE SYSTEM CASE STUDY: RESIDENT MICHAEL SCOTT (APT 4B)"
    bp1.font.name = "Segoe UI"
    bp1.font.size = Pt(11)
    bp1.font.bold = True
    bp1.font.color.rgb = RGBColor(56, 189, 248) # Sky-400

    bp2 = btf.add_paragraph()
    bp2.text = (
        "1. Input: Monthly meter logged 48.50 kL for October 2026 (Past baseline: 14.5–15.0 kL/month for 4 residents).\n"
        "2. ML Inference: LightGBM predicted expected bounds: Q10 = 13.0 kL | Q50 = 14.8 kL | Q95 = 16.8 kL upper threshold.\n"
        "3. Anomaly Decision: 48.50 kL is 3.2x baseline (exceeds Q95 ceiling) -> Risk Score = 0.96 (96% High Risk).\n"
        "4. UI Attribution: Status = POTENTIAL_LEAK | Pattern = SUDDEN_SPIKE | Estimated Extra Cost = ₹2,457 | Feedback = Unacknowledged."
    )
    bp2.font.name = "Segoe UI"
    bp2.font.size = Pt(9.5)
    bp2.font.color.rgb = RGBColor(226, 232, 240) # Slate-200

    output_path = "AquaFlow_ML_Pipeline_Slide.pptx"
    prs.save(output_path)
    print(f"Standalone ML pipeline presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_ml_pipeline_presentation()
